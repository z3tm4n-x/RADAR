"""Normalize stopping, range and LET workbook sheets to compact RADAR CSV tables."""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

ALUMINIUM_DENSITY_G_CM3 = 2.70

HZE_CHARGE_NUMBER_BY_SYMBOL: Mapping[str, int] = {
    "He": 2,
    "C": 6,
    "N": 7,
    "O": 8,
    "Ne": 10,
    "Mg": 12,
    "Si": 14,
    "S": 16,
    "Ar": 18,
    "Ca": 20,
    "Cr": 24,
    "Mn": 25,
    "Fe": 26,
    "Co": 27,
    "Ni": 28,
}

HZE_MASS_NUMBER_BY_SYMBOL: Mapping[str, float] = {
    "He": 4.0,
    "C": 12.0,
    "N": 14.0,
    "O": 16.0,
    "Ne": 20.0,
    "Mg": 24.0,
    "Si": 28.0,
    "S": 32.0,
    "Ar": 40.0,
    "Ca": 40.0,
    "Cr": 52.0,
    "Mn": 55.0,
    "Fe": 56.0,
    "Co": 59.0,
    "Ni": 58.0,
}

HZE_SYMBOLS: tuple[str, ...] = tuple(HZE_CHARGE_NUMBER_BY_SYMBOL)


@dataclass(frozen=True)
class StoppingTableCsvOutput:
    """CSV paths written by stopping workbook normalization."""

    proton_al_range: Path
    proton_si_let: Path
    hze_al_range: Path
    hze_si_let: Path


def _as_float(value: object) -> float | None:
    if value is None:
        return None

    if isinstance(value, str):
        text = value.strip().replace(",", ".")
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None

    if isinstance(value, int | float):
        return float(value)

    return None


def _header_name(value: object, fallback: str) -> str:
    if value is None:
        return fallback

    text = str(value).strip()
    if not text:
        return fallback

    return text


def _normalized_header_key(value: str) -> str:
    return value.strip().lower().replace(" ", "_").replace("-", "_")


def _find_header(headers: Sequence[str], candidates: Sequence[str]) -> str | None:
    normalized_to_original = {_normalized_header_key(header): header for header in headers}

    for candidate in candidates:
        normalized = _normalized_header_key(candidate)
        if normalized in normalized_to_original:
            return normalized_to_original[normalized]

    return None


def _worksheet_numeric_rows(
    workbook_path: str | Path,
    *,
    sheet_name: str,
) -> tuple[dict[str, float], ...]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    workbook = load_workbook(
        filename=str(workbook_path),
        data_only=True,
        read_only=True,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            msg = f"Sheet {sheet_name!r} not found. Available sheets: {workbook.sheetnames}"
            raise ValueError(msg)

        worksheet: Any = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)

        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            msg = f"Sheet {sheet_name!r} is empty."
            raise ValueError(msg) from exc

        headers = [
            _header_name(value, f"column_{index + 1}")
            for index, value in enumerate(raw_headers)
        ]

        if headers:
            headers[0] = "E"

        numeric_rows: list[dict[str, float]] = []

        for raw_row in rows:
            row: dict[str, float] = {}

            for header, value in zip(headers, raw_row, strict=False):
                numeric_value = _as_float(value)
                if numeric_value is not None:
                    row[header] = numeric_value

            if "E" in row:
                numeric_rows.append(row)

        numeric_rows.sort(key=lambda row: row["E"])
        return tuple(numeric_rows)
    finally:
        workbook.close()


def _require_column(rows: Sequence[dict[str, float]], column: str, sheet_name: str) -> None:
    if not rows:
        msg = f"Sheet {sheet_name!r} contains no numeric rows."
        raise ValueError(msg)

    if all(column not in row for row in rows):
        msg = f"Sheet {sheet_name!r} must contain numeric column {column!r}."
        raise ValueError(msg)


def _positive(value: float | None) -> bool:
    return value is not None and value > 0.0


def _require_positive_float(value: float | None, *, name: str) -> float:
    if value is None or value <= 0.0:
        msg = f"{name} must be a positive finite numeric value."
        raise ValueError(msg)

    return value


def _write_proton_al_range_csv(
    *,
    workbook_path: str | Path,
    output_path: Path,
    sheet_name: str,
    rho_al_g_cm3: float,
) -> None:
    rows = _worksheet_numeric_rows(workbook_path, sheet_name=sheet_name)

    _require_column(rows, "E", sheet_name)
    _require_column(rows, "H", sheet_name)
    _require_column(rows, "R", sheet_name)

    with output_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(
            (
                "energy_mev",
                "stopping_mev_cm2_g",
                "range_g_cm2",
                "range_mm",
                "source_sheet",
            )
        )

        for row in rows:
            energy_mev = row.get("E")
            stopping_mev_cm2_mg = row.get("H")
            range_mm = row.get("R")

            if not (_positive(energy_mev) and _positive(stopping_mev_cm2_mg) and _positive(range_mm)):
                continue

            energy_value = _require_positive_float(energy_mev, name="proton Al energy")
            stopping_value = _require_positive_float(
                stopping_mev_cm2_mg,
                name="proton Al stopping",
            )
            range_value = _require_positive_float(range_mm, name="proton Al range")

            writer.writerow(
                (
                    energy_value,
                    stopping_value * 1000.0,
                    rho_al_g_cm3 * range_value / 10.0,
                    range_value,
                    sheet_name,
                )
            )


def _write_proton_si_let_csv(
    *,
    workbook_path: str | Path,
    output_path: Path,
    sheet_name: str,
) -> None:
    rows = _worksheet_numeric_rows(workbook_path, sheet_name=sheet_name)

    if not rows:
        msg = f"Sheet {sheet_name!r} contains no numeric rows."
        raise ValueError(msg)

    headers = tuple(rows[0])
    electronic_column = _find_header(headers, ("H_coll", "H", "h_coll", "h"))
    nuclear_column = _find_header(headers, ("H_nucl", "H_n", "h_nucl", "h_n", "nuclear"))

    if electronic_column is None:
        msg = f"Sheet {sheet_name!r} must contain proton LET column H_coll or H."
        raise ValueError(msg)

    with output_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(
            (
                "energy_mev",
                "let_electronic_mev_cm2_mg",
                "let_nuclear_mev_cm2_mg",
                "let_total_mev_cm2_mg",
                "electronic_column",
                "nuclear_column",
                "source_sheet",
            )
        )

        for row in rows:
            energy_mev = row.get("E")
            let_electronic = row.get(electronic_column)
            let_nuclear = row.get(nuclear_column) if nuclear_column is not None else 0.0

            if not (_positive(energy_mev) and _positive(let_electronic)):
                continue

            energy_value = _require_positive_float(energy_mev, name="proton Si energy")
            let_electronic_value = _require_positive_float(
                let_electronic,
                name="proton Si electronic LET",
            )

            let_nuclear_value = 0.0
            if let_nuclear is not None and let_nuclear >= 0.0:
                let_nuclear_value = let_nuclear

            writer.writerow(
                (
                    energy_value,
                    let_electronic_value,
                    let_nuclear_value,
                    let_electronic_value + let_nuclear_value,
                    electronic_column,
                    nuclear_column or "",
                    sheet_name,
                )
            )


def _write_hze_al_range_csv(
    *,
    workbook_path: str | Path,
    output_path: Path,
    sheet_name: str,
    rho_al_g_cm3: float,
) -> None:
    rows = _worksheet_numeric_rows(workbook_path, sheet_name=sheet_name)
    _require_column(rows, "E", sheet_name)

    with output_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(
            (
                "z",
                "symbol",
                "mass_number",
                "mass_to_charge",
                "energy_mev_per_nucleon",
                "range_g_cm2",
                "range_mm",
                "source_sheet",
            )
        )

        for symbol in HZE_SYMBOLS:
            if all(symbol not in row for row in rows):
                continue

            charge_number = HZE_CHARGE_NUMBER_BY_SYMBOL[symbol]
            mass_number = HZE_MASS_NUMBER_BY_SYMBOL[symbol]
            mass_to_charge = mass_number / charge_number

            for row in rows:
                total_energy_mev = row.get("E")
                range_mm = row.get(symbol)

                if not (_positive(total_energy_mev) and _positive(range_mm)):
                    continue

                total_energy_value = _require_positive_float(
                    total_energy_mev,
                    name="HZE Al total energy",
                )
                range_value = _require_positive_float(range_mm, name="HZE Al range")

                writer.writerow(
                    (
                        charge_number,
                        symbol,
                        mass_number,
                        mass_to_charge,
                        total_energy_value / mass_number,
                        rho_al_g_cm3 * range_value / 10.0,
                        range_value,
                        sheet_name,
                    )
                )


def _rows_by_energy(rows: Sequence[dict[str, float]]) -> dict[float, dict[str, float]]:
    return {row["E"]: row for row in rows if "E" in row}


def _write_hze_si_let_csv(
    *,
    workbook_path: str | Path,
    output_path: Path,
    electronic_sheet_name: str,
    nuclear_sheet_name: str,
) -> None:
    electronic_rows = _worksheet_numeric_rows(workbook_path, sheet_name=electronic_sheet_name)

    try:
        nuclear_rows = _worksheet_numeric_rows(workbook_path, sheet_name=nuclear_sheet_name)
    except ValueError:
        nuclear_rows = ()

    _require_column(electronic_rows, "E", electronic_sheet_name)
    nuclear_by_energy = _rows_by_energy(nuclear_rows)

    with output_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(
            (
                "z",
                "symbol",
                "mass_number",
                "mass_to_charge",
                "energy_mev_per_nucleon",
                "let_electronic_mev_cm2_mg",
                "let_nuclear_mev_cm2_mg",
                "let_total_mev_cm2_mg",
                "electronic_source_sheet",
                "nuclear_source_sheet",
            )
        )

        for symbol in HZE_SYMBOLS:
            if all(symbol not in row for row in electronic_rows):
                continue

            charge_number = HZE_CHARGE_NUMBER_BY_SYMBOL[symbol]
            mass_number = HZE_MASS_NUMBER_BY_SYMBOL[symbol]
            mass_to_charge = mass_number / charge_number

            for row in electronic_rows:
                total_energy_mev = row.get("E")
                let_electronic = row.get(symbol)

                if not (_positive(total_energy_mev) and _positive(let_electronic)):
                    continue

                total_energy_value = _require_positive_float(
                    total_energy_mev,
                    name="HZE Si total energy",
                )
                let_electronic_value = _require_positive_float(
                    let_electronic,
                    name="HZE Si electronic LET",
                )

                nuclear_row = nuclear_by_energy.get(total_energy_value)
                let_nuclear = nuclear_row.get(symbol) if nuclear_row is not None else 0.0

                let_nuclear_value = 0.0
                if let_nuclear is not None and let_nuclear >= 0.0:
                    let_nuclear_value = let_nuclear

                writer.writerow(
                    (
                        charge_number,
                        symbol,
                        mass_number,
                        mass_to_charge,
                        total_energy_value / mass_number,
                        let_electronic_value,
                        let_nuclear_value,
                        let_electronic_value + let_nuclear_value,
                        electronic_sheet_name,
                        nuclear_sheet_name if nuclear_rows else "",
                    )
                )


def convert_stopping_workbook_to_csvs(
    *,
    workbook_path: str | Path,
    output_dir: str | Path,
    rho_al_g_cm3: float = ALUMINIUM_DENSITY_G_CM3,
    proton_al_sheet: str = "p_Al",
    proton_si_sheet: str = "p_Si",
    hze_al_sheet: str = "i_R_Al",
    hze_si_e_sheet: str = "i_Si_e",
    hze_si_n_sheet: str = "i_Si_n",
) -> StoppingTableCsvOutput:
    """Convert workbook stopping/range/LET sheets to RADAR normative CSV tables."""

    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)

    output = StoppingTableCsvOutput(
        proton_al_range=destination / "proton_al_range.csv",
        proton_si_let=destination / "proton_si_let.csv",
        hze_al_range=destination / "hze_al_range.csv",
        hze_si_let=destination / "hze_si_let.csv",
    )

    _write_proton_al_range_csv(
        workbook_path=workbook_path,
        output_path=output.proton_al_range,
        sheet_name=proton_al_sheet,
        rho_al_g_cm3=rho_al_g_cm3,
    )
    _write_proton_si_let_csv(
        workbook_path=workbook_path,
        output_path=output.proton_si_let,
        sheet_name=proton_si_sheet,
    )
    _write_hze_al_range_csv(
        workbook_path=workbook_path,
        output_path=output.hze_al_range,
        sheet_name=hze_al_sheet,
        rho_al_g_cm3=rho_al_g_cm3,
    )
    _write_hze_si_let_csv(
        workbook_path=workbook_path,
        output_path=output.hze_si_let,
        electronic_sheet_name=hze_si_e_sheet,
        nuclear_sheet_name=hze_si_n_sheet,
    )

    return output


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Convert RADAR stopping/range/LET workbook sheets to compact CSV tables."
    )
    parser.add_argument("workbook_path")
    parser.add_argument("output_dir")
    parser.add_argument("--rho-al", type=float, default=ALUMINIUM_DENSITY_G_CM3)

    args = parser.parse_args(argv)

    output = convert_stopping_workbook_to_csvs(
        workbook_path=args.workbook_path,
        output_dir=args.output_dir,
        rho_al_g_cm3=args.rho_al,
    )

    print(f"proton Al range: {output.proton_al_range}")
    print(f"proton Si LET:   {output.proton_si_let}")
    print(f"HZE Al range:    {output.hze_al_range}")
    print(f"HZE Si LET:      {output.hze_si_let}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
