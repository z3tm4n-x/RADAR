import csv
from pathlib import Path

import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]

from radar.shielding.stopping_tables import convert_stopping_workbook_to_csvs


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as file:
        return list(csv.DictReader(file))


def _sample_workbook(path: Path) -> None:
    workbook = Workbook()

    p_al = workbook.active
    p_al.title = "p_Al"
    p_al.append(("E", "H", "R"))
    p_al.append((1.0, 2.0, 0.1))
    p_al.append((10.0, 3.0, 1.0))

    e_al = workbook.create_sheet("e_Al")
    e_al.append(("E", "e", "R"))
    e_al.append((0.1, 0.01, 2.0))
    e_al.append((1.0, 0.02, 20.0))

    p_si = workbook.create_sheet("p_Si")
    p_si.append(("E", "H_coll", "H_nucl"))
    p_si.append((1.0, 4.0, 0.1))
    p_si.append((10.0, 5.0, 0.2))

    i_r_al = workbook.create_sheet("i_R_Al")
    i_r_al.append(("E", "He", "C"))
    i_r_al.append((40.0, 0.5, 1.0))
    i_r_al.append((400.0, 5.0, 10.0))

    i_si_e = workbook.create_sheet("i_Si_e")
    i_si_e.append(("E", "He", "C"))
    i_si_e.append((40.0, 6.0, 7.0))
    i_si_e.append((400.0, 8.0, 9.0))

    i_si_n = workbook.create_sheet("i_Si_n")
    i_si_n.append(("E", "He", "C"))
    i_si_n.append((40.0, 0.6, 0.7))
    i_si_n.append((400.0, 0.8, 0.9))

    workbook.save(path)


def test_convert_stopping_workbook_to_csvs(tmp_path: Path) -> None:
    workbook_path = tmp_path / "tables.xlsx"
    output_dir = tmp_path / "out"
    _sample_workbook(workbook_path)

    output = convert_stopping_workbook_to_csvs(
        workbook_path=workbook_path,
        output_dir=output_dir,
    )

    assert output.proton_al_range.exists()
    assert output.electron_al_range.exists()
    assert output.proton_si_let.exists()
    assert output.hze_al_range.exists()
    assert output.hze_si_let.exists()

    proton_al = _read_csv(output.proton_al_range)
    assert proton_al[0]["energy_mev"] == "1.0"
    assert float(proton_al[0]["stopping_mev_cm2_g"]) == pytest.approx(2000.0)
    assert float(proton_al[0]["range_g_cm2"]) == pytest.approx(0.027)

    electron_al = _read_csv(output.electron_al_range)
    assert electron_al[0]["energy_mev"] == "0.1"
    assert float(electron_al[0]["stopping_mev_cm2_g"]) == pytest.approx(10.0)
    assert float(electron_al[0]["range_g_cm2"]) == pytest.approx(0.002)
    assert float(electron_al[0]["range_mg_cm2"]) == pytest.approx(2.0)

    proton_si = _read_csv(output.proton_si_let)
    assert float(proton_si[0]["let_electronic_mev_cm2_mg"]) == pytest.approx(4.0)
    assert float(proton_si[0]["let_nuclear_mev_cm2_mg"]) == pytest.approx(0.1)
    assert float(proton_si[0]["let_total_mev_cm2_mg"]) == pytest.approx(4.1)

    hze_al = _read_csv(output.hze_al_range)
    he_rows = [row for row in hze_al if row["symbol"] == "He"]
    assert float(he_rows[0]["energy_mev_per_nucleon"]) == pytest.approx(10.0)
    assert float(he_rows[0]["range_g_cm2"]) == pytest.approx(0.135)

    hze_si = _read_csv(output.hze_si_let)
    c_rows = [row for row in hze_si if row["symbol"] == "C"]
    assert float(c_rows[0]["energy_mev_per_nucleon"]) == pytest.approx(40.0 / 12.0)
    assert float(c_rows[0]["let_total_mev_cm2_mg"]) == pytest.approx(7.7)


def test_convert_stopping_workbook_rejects_missing_sheet(tmp_path: Path) -> None:
    workbook_path = tmp_path / "tables.xlsx"
    _sample_workbook(workbook_path)

    with pytest.raises(ValueError, match="not found"):
        convert_stopping_workbook_to_csvs(
            workbook_path=workbook_path,
            output_dir=tmp_path / "out",
            proton_al_sheet="missing",
        )
